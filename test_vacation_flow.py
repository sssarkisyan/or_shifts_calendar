"""Скриптовая проверка портала отпусков — не часть приложения, только для
разработки. Прогоняет: подсчёт дней с праздниками, правила (лимит дней 31 +
остаток, часть ≥14 дней, пересечение основных ≤7 дн., не больше 2 основных
одновременно — блокируют только одобренные отпуска коллег, поданные лишь
предупреждают), подачу/удаление через UI, календарь, согласование, Excel."""
import datetime
import io
import json
import os

os.environ["VACATION_LOCAL_BACKEND"] = "1"  # никогда не трогать настоящую Google Таблицу

from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

import config
import vacation_config as vc
from secrets_util import get_secret
import streamlit as st_lib
from vacation_backend import BACKEND_NAME, LOCAL_FILE, LOCAL_LOG_FILE, backend, new_record
from vacation_logic import (APPROVED, PENDING, REJECTED, TeamRules, build_excel, editor_updates, vacation_days,
                            validate_new_period)

assert BACKEND_NAME == "local_csv"
for f in (LOCAL_FILE, LOCAL_LOG_FILE):
    if os.path.exists(f):
        os.remove(f)

errors = []


def check(cond, msg):
    if not cond:
        errors.append(msg)
    print(("OK  " if cond else "FAIL"), msg)


D = datetime.date
Y = vc.PLAN_YEAR
names = [s["name"] for s in config.STAFF]
me = names[3]

# --- подсчёт дней ---
check(vacation_days(D(2027, 1, 1), D(2027, 1, 14)) == 6, "1–14 января 2027: 14 кал. дней − 8 праздников = 6")
check(vacation_days(D(2027, 7, 1), D(2027, 7, 14)) == 14, "июль без праздников: 14 дней")
check(vacation_days(D(2027, 6, 7), D(2027, 6, 20)) == 13, "12 июня внутри периода не считается")


def rec(name, a, b, status=PENDING):
    return new_record(name, a, b, "", status, "t")


MAIN = frozenset(st["name"] for st in config.STAFF if st["kind"] in vc.OVERLAP_KINDS)
RULES = TeamRules(MAIN, vc.MAX_OVERLAP_DAYS, vc.MAX_CONCURRENT_MAIN)
part_timer = [st["name"] for st in config.STAFF if st["kind"] == "совместитель"][0]
colleague, colleague2, colleague3, colleague4 = names[4], names[5], names[6], names[7]
check({me, colleague, colleague2, colleague3, colleague4} <= MAIN and part_timer not in MAIN,
      "sanity: пятеро основных и совместитель")


def check_all(name, a, b, records, ent=None):
    return validate_new_period(name, a, b, Y, records, vc.entitlement(name) if ent is None else ent,
                               RULES, vc.MIN_LONG_PART)


def validate(*args, **kw):
    return check_all(*args, **kw)[0]


def warn(*args, **kw):
    return check_all(*args, **kw)[1]


# --- лимит дней: 31 + остаток ---
check(vc.entitlement(me) == 31, "положено 31 день по умолчанию")
check(validate(me, D(Y, 7, 1), D(Y, 7, 31), []) == [], "31 день — ок")
check(any("лимит" in e for e in validate(me, D(Y, 7, 1), D(Y, 8, 1), [])), "32 дня — блок")
vc.CARRYOVER_DAYS[me] = 5
check(vc.entitlement(me) == 36, "остаток 5 дней → положено 36")
check(validate(me, D(Y, 7, 1), D(Y, 8, 5), []) == [], "36 дней с остатком — ок")
check(any("лимит" in e for e in validate(me, D(Y, 7, 1), D(Y, 8, 6), [])), "37 дней с остатком — блок")
del vc.CARRYOVER_DAYS[me]

# --- часть ≥14 дней и свои пересечения ---
existing = [rec(me, D(Y, 3, 1), D(Y, 3, 10))]
check(any("ст. 125" in e for e in validate(me, D(Y, 7, 1), D(Y, 7, 10), existing)),
      "10+10 дней без части ≥14, осталось 11 — блок")
check(validate(me, D(Y, 7, 1), D(Y, 7, 7), []) == [], "первая короткая часть 7 дней, осталось 24 — ок")
check(any("Пересекается" in e for e in validate(me, D(Y, 3, 5), D(Y, 3, 20), existing)), "пересечение со своим периодом — блок")
check(validate(me, D(Y, 3, 5), D(Y, 3, 20), [rec(me, D(Y, 3, 1), D(Y, 3, 10), REJECTED)]) == [],
      "отклонённый период не мешает")
check(any("%d году" % Y in e for e in validate(me, D(Y, 12, 25), D(Y + 1, 1, 10), [])), "период через Новый год — блок")

# --- поданные (не рассмотренные) заявки коллег НЕ блокируют, только предупреждают ---
crowd = [rec(colleague, D(Y, 7, 1), D(Y, 7, 14)), rec(colleague2, D(Y, 7, 1), D(Y, 7, 14))]
check(validate(me, D(Y, 7, 1), D(Y, 7, 14), crowd) == [], "двое коллег подали те же даты — подать можно")
w = warn(me, D(Y, 7, 1), D(Y, 7, 14), crowd)
check(any(colleague in x and "14 дн." in x for x in w) and any("максимум" in x for x in w),
      "предупреждение: пересечение 14 дн. и уже двое на эти даты")
check(warn(me, D(Y, 8, 1), D(Y, 8, 14), crowd) == [], "свободные даты — без предупреждений")

# --- одобренные отпуска блокируют: пересечение > 7 дней ---
theirs = [rec(colleague, D(Y, 7, 1), D(Y, 7, 14), APPROVED)]
check(validate(me, D(Y, 7, 8), D(Y, 7, 21), theirs) == [], "с одобренным пересечение ровно 7 дней — ок")
errs = validate(me, D(Y, 7, 7), D(Y, 7, 20), theirs)
check(any(colleague in e and "8 дн." in e and "одобрены" in e for e in errs), "с одобренным 8 дней — блок")
check(any(colleague in e for e in validate(me, D(Y, 6, 20), D(Y, 7, 31), theirs)), "накрывает одобренный — блок")
theirs[0]["status"] = REJECTED
check(validate(me, D(Y, 7, 1), D(Y, 7, 14), theirs) == [], "отклонённый отпуск коллеги не мешает")
theirs[0]["status"] = APPROVED
check(validate(part_timer, D(Y, 7, 1), D(Y, 7, 14), theirs) == [], "совместитель — пересекаться можно полностью")
check(validate(me, D(Y, 7, 1), D(Y, 7, 14), [rec(part_timer, D(Y, 7, 1), D(Y, 7, 14), APPROVED)]) == [],
      "основной с совместителем — пересекаться можно полностью")

# --- одобренные: не больше 2 основных в один день ---
check(vc.MAX_CONCURRENT_MAIN == 2, "лимит одновременных основных = 2")
pair = [rec(colleague, D(Y, 7, 1), D(Y, 7, 7), APPROVED), rec(colleague2, D(Y, 7, 5), D(Y, 7, 11), APPROVED)]
errs = validate(me, D(Y, 7, 7), D(Y, 7, 20), pair)
check(len(errs) == 1 and "максимум" in errs[0] and "07.07" in errs[0], "7 июля одобрено двоим — недоступно")
check(validate(me, D(Y, 7, 8), D(Y, 7, 21), pair) == [], "с 8 июля одобрен только один — ок")
check(validate(part_timer, D(Y, 7, 1), D(Y, 7, 14), pair) == [], "совместитель лимитом не ограничен")
pair[1]["status"] = PENDING
check(validate(me, D(Y, 7, 7), D(Y, 7, 20), pair) == [], "если вторая заявка ещё не одобрена — подать можно")

# --- UI: сотрудник ---
at = AppTest.from_file("vacation_app.py", default_timeout=30)
at.run()
check(not at.exception, "загрузка без ошибок")
check(any("ЛОКАЛЬНЫЙ РЕЖИМ" in e.value for e in at.error), "приложение в тесте работает с локальным файлом")


def as_employee(n):
    at.selectbox(key="me_name").select(n).run()


def pick(a, b):
    key = [k for k in at.session_state.filtered_state if str(k).startswith("range_")][0]
    at.date_input(key=key).set_value((a, b)).run()


def refresh():
    """Тест пишет в хранилище напрямую — сбросить кэш приложения и перерисовать."""
    st_lib.cache_data.clear()
    at.run()


def confirm():
    key = [k for k in at.session_state.filtered_state if str(k).startswith("confirm_")][0]
    at.checkbox(key=key).check().run()


def log_actions(name):
    return [e["action"] for e in backend.list_log(Y) if e["name"] == name]


def submit_btn():
    return [b for b in at.button if b.label == "Подать заявку"][0]


def calendar_html():
    return at.get("html")[0].proto.body


as_employee(me)
pick(D(Y, 7, 1), D(Y, 7, 14))
check(submit_btn().disabled, "без подтверждения «подаю от своего имени» кнопка неактивна")
check(any(me in c.label for c in at.checkbox), "в подтверждении указано ФИО")
confirm()
check(not submit_btn().disabled, "после подтверждения кнопка активна")
submit_btn().click().run()
check(not at.exception, "подача без ошибок")
mine = [r for r in backend.list_vacations(Y) if r["name"] == me]
check(len(mine) == 1 and mine[0]["days"] == 14 and mine[0]["status"] == PENDING, "заявка сохранена, 14 дн., на рассмотрении")
check(any("Заявка подана" in s.value for s in at.success), "сообщение об успехе")
check(log_actions(me) == ["Подана заявка"], "журнал: «Подана заявка»")
check(not at.checkbox or not at.checkbox[0].value, "после подачи подтверждение сброшено")

pick(D(Y, 7, 10), D(Y, 7, 20))
check(any("Пересекается" in e.value for e in at.error) and submit_btn().disabled,
      "своё пересечение — ошибка видна сразу, кнопка неактивна")

pick(D(Y, 10, 1), D(Y, 10, 7))
confirm()
submit_btn().click().run()
second = [r for r in backend.list_vacations(Y) if r["name"] == me and r["start"].month == 10]
check(len(second) == 1, "вторая заявка сохранена")
[b for b in at.button if b.key == "del_" + second[0]["id"]][0].click().run()
check(len([r for r in backend.list_vacations(Y) if r["name"] == me]) == 1, "удаление своей заявки (через подтверждение)")
check(log_actions(me) == ["Подана заявка", "Подана заявка", "Заявка удалена"], "журнал: подача, подача, удаление")
check(any("История действий" in e.label for e in at.expander), "сотрудник видит историю своих заявок")

# коллега подаёт те же даты — можно, с предупреждением
as_employee(colleague)
html_text = calendar_html()
check('class="vc-p' in html_text and me in html_text, "календарь: заявка коллеги видна бледно-розовым, с фамилией")
pick(D(Y, 7, 1), D(Y, 7, 14))
confirm()
check(any(me in w.value for w in at.warning) and not submit_btn().disabled,
      "предупреждение о заявке коллеги, но подать можно")
submit_btn().click().run()
check(len([r for r in backend.list_vacations(Y) if r["name"] == colleague]) == 1, "заявка на те же даты сохранена")

# --- UI: руководитель ---
backend.add_vacation(rec(colleague3, D(Y, 10, 1), D(Y, 10, 14)))  # ни с кем не пересекается
st_lib.cache_data.clear()
at.text_input(key="mgr_pw_input").input(get_secret("manager_passcode", config.MANAGER_PASSCODE_FALLBACK)).run()
[b for b in at.button if b.key == "mgr_login_btn"][0].click().run()
check(at.session_state["manager_unlocked"], "вход руководителя")
check(not at.exception, "свод без ошибок")
check(not any("ОДОБРЕННЫХ" in e.value for e in at.error), "пока ничего не одобрено — нарушений нет")
approve = [b for b in at.button if b.label.startswith("Одобрить без конфликтов")]
check(len(approve) == 1 and "(1)" in approve[0].label, "«Одобрить без конфликтов» — только 1 заявка")
approve[0].click().run()
st_by_name = {r["name"]: r["status"] for r in backend.list_vacations(Y)}
check(st_by_name[colleague3] == APPROVED and st_by_name[me] == PENDING and st_by_name[colleague] == PENDING,
      "одобрена только бесконфликтная, спорные остались на рассмотрении")
check(log_actions(colleague3) == ["Статус: Одобрено"], "журнал: одобрение записано")
check([e["device"] for e in backend.list_log(Y) if e["name"] == colleague3] == ["руководитель"], "журнал: кто — руководитель")
check(any("Журнал действий" in e.label for e in at.expander), "руководитель видит журнал")

# руководитель одобряет «меня», коллегу — отклоняет
ids = {r["name"]: r["id"] for r in backend.list_vacations(Y)}
backend.update_vacations({ids[me]: {"status": APPROVED, "manager_comment": ""},
                          ids[colleague]: {"status": REJECTED, "manager_comment": "перенесите на август"}}, "t2")
backend.add_vacation(rec(colleague2, D(Y, 7, 3), D(Y, 7, 9), APPROVED))  # 7 дн. с «мной» — допустимо
refresh()
check(not at.exception and not any("ОДОБРЕННЫХ" in e.value for e in at.error), "одобренные без нарушений")

# 1–9 июля... 3–9 июля одобрено двоим → для остальных основных недоступно
as_employee(colleague4)
html_text = calendar_html()
check('class="vc-full' in html_text, "календарь: дни с двумя одобренными — серые/недоступные")
pick(D(Y, 7, 5), D(Y, 7, 20))
check(any("недоступны" in e.value for e in at.error) and submit_btn().disabled, "3-й основной на 05.07 — блок")
pick(D(Y, 7, 10), D(Y, 7, 23))
check(not at.error or all("ЛОКАЛЬНЫЙ" in e.value for e in at.error), "с 10 июля (одобрен один, пересечение 5 дн.) — ок")
as_employee(part_timer)
check('class="vc-full' not in calendar_html(), "у совместителя недоступных дней нет")

# заведующий и профессор видят заявки друг друга; основные их не видят, лимиты не задеты
prof = [x["name"] for x in config.STAFF if x["kind"] == "профессор"][0]
head = [x["name"] for x in config.STAFF if x["kind"] == "заведующий"][0]
backend.add_vacation(rec(prof, D(Y, 3, 1), D(Y, 3, 14)))
refresh()
as_employee(head)
check(prof in calendar_html() and 'class="vc-p' in calendar_html(), "заведующий видит заявку профессора")
check(me not in calendar_html(), "заведующий не видит отпуска основных врачей")
as_employee(me)
check(prof not in calendar_html(), "основной врач не видит заявку профессора")
pick(D(Y, 3, 1), D(Y, 3, 14))
check(not any("поданы заявки" in w.value for w in at.warning), "отпуск профессора не создаёт конфликтов для основных")
as_employee(part_timer)
check(me in calendar_html(), "совместитель по-прежнему видит отпуска основных")

# нарушение среди одобренных видно руководителю
backend.add_vacation(rec(colleague4, D(Y, 7, 5), D(Y, 7, 6), APPROVED))
refresh()
check(any("ОДОБРЕННЫХ" in e.value and "05.07–06.07" in e.value for e in at.error),
      "руководитель видит 3 одобренных основных 5–6 июля")
htmls = at.get("html")
check(len(htmls) == 2, "календарь есть и у сотрудника, и в своде")
mgr_html = htmls[1].proto.body
check('class="vc-bad' in mgr_html, "календарь свода: нарушение среди одобренных — красным")
check('class="vc-full' in mgr_html, "календарь свода: закрытые для основных дни — серым")
check(prof + " (профессор)" in mgr_html and me in mgr_html, "календарь свода: видны все сотрудники, с типом у неосновных")

# правка в таблице согласования → только изменённые строки уходят в сохранение
import pandas as pd
orig = pd.DataFrame({"Статус": ["Одобрено", "Одобрено"], "Комментарий руководителя": ["", ""]}, index=["a", "b"])
edit = orig.copy()
edit.at["b", "Статус"] = "Отклонено"
edit.at["b", "Комментарий руководителя"] = "пересекается"
check(editor_updates(orig, edit) == {"b": {"status": REJECTED, "manager_comment": "пересекается"}},
      "изменения из таблицы согласования определяются верно")
edit2 = orig.copy()
edit2.at["a", "Комментарий руководителя"] = None
check(editor_updates(orig, edit2) == {}, "очищенная пустая ячейка не считается изменением")

# --- Excel ---
recs = backend.list_vacations(Y)
xlsx = build_excel(Y, config.STAFF, recs, vc.entitlement, vc.carryover, RULES, vc.MIN_LONG_PART)
wb = load_workbook(io.BytesIO(xlsx))
check(wb.sheetnames[0] == "Свод" and "Июль" in wb.sheetnames and "Заявки" in wb.sheetnames, "листы Excel")
jul = wb["Июль"]
row = [r for r in jul.iter_rows(values_only=True) if r[0] == me][0]
check(row[1:15] == ("От",) * 14 and row[15] is None, "календарь июля: 1–14 «От»")
check([r for r in jul.iter_rows(values_only=True) if r[0] == colleague][0][1] is None, "отклонённый в календарь не попал")
appr_row = [r for r in jul.iter_rows() if r[0].value == "Основных: одобрено"][0]
check(appr_row[5].value == 3 and appr_row[5].fill.fgColor.rgb.endswith("FF7C80"), "5 июля: 3 одобренных, красным")
check(appr_row[1].value == 1 and appr_row[1].fill.fgColor.rgb in (None, "00000000"), "1 июля: 1 одобренный, без подсветки")

os.remove(LOCAL_FILE)
os.remove(LOCAL_LOG_FILE)
print()
print("ВСЁ ОК" if not errors else "ОШИБОК: %d" % len(errors))
