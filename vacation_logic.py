"""Подсчёт дней отпуска, проверки правил и выгрузка в Excel — без Streamlit,
чтобы это можно было проверять отдельно (test_vacation_flow.py)."""
import calendar as calmod
import dataclasses
import datetime
import functools
import io

import holidays as holidays_lib
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PENDING, APPROVED, REJECTED = "PENDING", "APPROVED", "REJECTED"
STATUS_LABELS = {PENDING: "На рассмотрении", APPROVED: "Одобрено", REJECTED: "Отклонено"}
ACTIVE_STATUSES = {PENDING, APPROVED}  # отклонённые заявки дни не занимают

MONTH_NAMES = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль",
               "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]
DOW_RU = ["Пн", "Вт", "Ср", "Чт", "Пт", "Сб", "Вс"]


@functools.lru_cache(maxsize=None)
def ru_holidays(year):
    """Праздничные нерабочие дни РФ (ст. 112 ТК) — без переносов выходных."""
    return frozenset(holidays_lib.Russia(years=year).keys())


def is_holiday(d):
    return d in ru_holidays(d.year)


def daterange(start, end):
    d = start
    while d <= end:
        yield d
        d += datetime.timedelta(days=1)


def vacation_days(start, end):
    """Дней отпуска в периоде: календарные дни минус праздничные (ст. 120 ТК)."""
    return sum(1 for d in daterange(start, end) if not is_holiday(d))


def fmt_date(d):
    return d.strftime("%d.%m.%Y")


def fmt_period(start, end):
    return "%s – %s" % (fmt_date(start), fmt_date(end))


MONTHS_GEN = ["января", "февраля", "марта", "апреля", "мая", "июня", "июля",
              "августа", "сентября", "октября", "ноября", "декабря"]


def fmt_date_words(d):
    return "%d %s %d" % (d.day, MONTHS_GEN[d.month - 1], d.year)


def fmt_period_words(start, end):
    """«с 1 июля по 14 июля 2027» (год один раз, если совпадает)."""
    if start.year == end.year:
        return "с %d %s по %s" % (start.day, MONTHS_GEN[start.month - 1], fmt_date_words(end))
    return "с %s по %s" % (fmt_date_words(start), fmt_date_words(end))


def plural(n, one, few, many):
    """plural(5, "день", "дня", "дней") → "дней"."""
    if n % 10 == 1 and n % 100 != 11:
        return one
    if 2 <= n % 10 <= 4 and not 12 <= n % 100 <= 14:
        return few
    return many


def n_days(n):
    """5 → "5 дней", 1 → "1 день", 3 → "3 дня"."""
    return "%d %s" % (n, plural(n, "день", "дня", "дней"))


def n_doctors(n):
    return "%d %s" % (n, plural(n, "основной врач", "основных врача", "основных врачей"))


def active(records):
    return [r for r in records if r["status"] in ACTIVE_STATUSES]


def occupancy(records, exclude_name=None):
    """{дата: [ФИО, ...]} — кто в отпуске в этот день (активные заявки)."""
    out = {}
    for r in active(records):
        if r["name"] == exclude_name:
            continue
        for d in daterange(r["start"], r["end"]):
            out.setdefault(d, []).append(r["name"])
    return out


def collapse_dates(dates):
    """[даты] → [(начало, конец)] — подряд идущие дни склеиваются в интервалы."""
    ranges = []
    for d in sorted(dates):
        if ranges and d - ranges[-1][1] == datetime.timedelta(days=1):
            ranges[-1][1] = d
        else:
            ranges.append([d, d])
    return [(a, b) for a, b in ranges]


def fmt_ranges(dates):
    return ", ".join(
        fmt_date(a) if a == b else "%s–%s" % (a.strftime("%d.%m"), fmt_date(b))
        for a, b in collapse_dates(dates)
    )


@dataclasses.dataclass(frozen=True)
class TeamRules:
    """Правила для основных сотрудников (main_names): отпуска двух из них
    пересекаются не больше max_overlap календарных дней, и в один день в
    отпуске не больше max_concurrent из них. Совместители не ограничены.
    Блокируют подачу только ОДОБРЕННЫЕ отпуска коллег — поданные, но не
    рассмотренные заявки лишь показываются (кто первым подал — не важно,
    решает руководитель)."""
    main_names: frozenset
    max_overlap: int
    max_concurrent: int


def overlap_days(a, b):
    """Сколько календарных дней пересекаются два периода (заявки)."""
    return max(0, (min(a["end"], b["end"]) - max(a["start"], b["start"])).days + 1)


def team_conflicts(name, start, end, others, rules):
    """Нарушения командных правил, если сотрудник name уйдёт в отпуск
    start–end, а others — уже учитываемые отпуска коллег. Список строк."""
    if name not in rules.main_names:
        return []
    others = [r for r in others if r["name"] != name and r["name"] in rules.main_names]
    out = []
    new = {"start": start, "end": end}
    for r in sorted(others, key=lambda r: r["start"]):
        n = overlap_days(new, r)
        if n > rules.max_overlap:
            out.append("пересечение с %s (%s) — %s, допускается не больше %s"
                       % (r["name"], fmt_period(r["start"], r["end"]), n_days(n), n_days(rules.max_overlap)))
    occ = occupancy(others)
    full = [d for d in daterange(start, end) if len(occ.get(d, [])) >= rules.max_concurrent]
    if full:
        who = sorted({n for d in full for n in occ[d]})
        out.append("%s в отпуске уже %s (%s) — это максимум"
                   % (fmt_ranges(full), n_doctors(rules.max_concurrent), ", ".join(who)))
    return out


def approved(records):
    return [r for r in records if r["status"] == APPROVED]


def validate_new_period(name, start, end, year, all_records, entitlement, rules, min_long_part):
    """(ошибки, предупреждения). Ошибки блокируют подачу: свои правила (лимит
    дней, часть ≥14, пересечение со своим периодом) и конфликт с ОДОБРЕННЫМИ
    отпусками коллег. Предупреждения — конфликт с поданными заявками коллег:
    подать можно, решение за руководителем."""
    if start > end:
        return ["Дата окончания раньше даты начала."], []
    if start.year != year or end.year != year:
        return ["Отпуск должен целиком лежать в %d году." % year], []
    days = vacation_days(start, end)
    if days == 0:
        return ["В выбранных датах только праздничные дни — они в отпуск не входят."], []

    errors = []
    mine = [r for r in active(all_records) if r["name"] == name]
    for r in mine:
        if start <= r["end"] and r["start"] <= end:
            errors.append("Пересекается с вашим отпуском %s." % fmt_period(r["start"], r["end"]))

    used = sum(r["days"] for r in mine)
    if used + days > entitlement:
        errors.append(
            "Превышен лимит: положено %s, уже запланировано %s, в этом отпуске %s (осталось %s)."
            % (n_days(entitlement), n_days(used), n_days(days), n_days(max(entitlement - used, 0)))
        )
    else:
        has_long = any(r["days"] >= min_long_part for r in mine)
        remaining_after = entitlement - used - days
        if not has_long and days < min_long_part and remaining_after < min_long_part:
            errors.append(
                "Хотя бы одна часть отпуска должна быть не меньше %s (ст. 125 ТК РФ). "
                "После этого отпуска останется %s — длинную часть будет уже не набрать. "
                "Удлините отпуск или сначала запланируйте длинную часть."
                % (n_days(min_long_part), n_days(remaining_after))
            )

    errors += ["Даты недоступны: " + c + " (отпуска коллег уже одобрены)."
               for c in team_conflicts(name, start, end, approved(all_records), rules)]
    warnings = [] if errors else [
        "На эти даты уже поданы заявки: " + c + ". Подать можно — решение примет заведующий."
        for c in team_conflicts(name, start, end, active(all_records), rules)
    ]
    return errors, warnings


def record_conflicts(rec, records, rules):
    """Конфликты заявки rec с остальными активными заявками (для руководителя)."""
    others = [r for r in active(records) if r["id"] != rec["id"]]
    return team_conflicts(rec["name"], rec["start"], rec["end"], others, rules)


def approved_violations(records, rules):
    """Нарушения среди одобренных отпусков: (пары с пересечением > max_overlap,
    {дата: [ФИО]} где одобренных основных больше max_concurrent)."""
    appr = [r for r in approved(records) if r["name"] in rules.main_names]
    pairs = []
    for i, a in enumerate(appr):
        for b in appr[i + 1:]:
            if a["name"] != b["name"] and overlap_days(a, b) > rules.max_overlap:
                pairs.append((a, b, overlap_days(a, b)))
    crowded = {d: who for d, who in occupancy(appr).items() if len(who) > rules.max_concurrent}
    return pairs, crowded


def bad_days(records, rules):
    """Дни, попадающие в нарушения среди одобренных отпусков."""
    pairs, crowded = approved_violations(records, rules)
    out = set(crowded)
    for a, b, _ in pairs:
        out.update(daterange(max(a["start"], b["start"]), min(a["end"], b["end"])))
    return out


def team_calendar(records, names, exclude_name=None):
    """{дата: {"approved": [ФИО], "pending": [ФИО]}} — отпуска сотрудников из
    names (кроме exclude_name) по дням."""
    out = {}
    for r in active(records):
        if r["name"] == exclude_name or r["name"] not in names:
            continue
        key = "approved" if r["status"] == APPROVED else "pending"
        for d in daterange(r["start"], r["end"]):
            out.setdefault(d, {"approved": [], "pending": []})[key].append(r["name"])
    return out


def person_summary(name, records, entitlement, min_long_part):
    mine = [r for r in records if r["name"] == name]
    act = active(mine)
    return {
        "entitlement": entitlement,
        "planned": sum(r["days"] for r in act),
        "approved": sum(r["days"] for r in act if r["status"] == APPROVED),
        "pending": sum(r["days"] for r in act if r["status"] == PENDING),
        "has_long": any(r["days"] >= min_long_part for r in act),
        "n_periods": len(mine),
    }


def editor_updates(original, edited):
    """Сравнить таблицу согласования до/после правки руководителем (DataFrame,
    индекс — ID заявки). Возвращает {id: {"status": код, "manager_comment": ..}}
    только для изменившихся строк."""
    label_to_status = {v: k for k, v in STATUS_LABELS.items()}
    updates = {}
    for rec_id in edited.index:
        label = edited.at[rec_id, "Статус"]
        comment = edited.at[rec_id, "Комментарий заведующего"] or ""
        if (label, comment) != (original.at[rec_id, "Статус"], original.at[rec_id, "Комментарий заведующего"]):
            updates[rec_id] = {"status": label_to_status[label], "manager_comment": comment}
    return updates


# ---------- Excel ----------

# Те же цвета, что в приложении: бледно-зелёный — одобрено, бледно-розовый — подано.
FILL_APPROVED = PatternFill("solid", fgColor="CDEBD3")
FILL_PENDING = PatternFill("solid", fgColor="F9D5DC")
FILL_TODO = PatternFill("solid", fgColor="FFE699")
FILL_WEEKEND = PatternFill("solid", fgColor="E7E6E6")
FILL_HOLIDAY = PatternFill("solid", fgColor="F4B183")
FILL_BAD = PatternFill("solid", fgColor="FF7C80")
BOLD = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center")


def build_excel(year, staff, records, entitlement_fn, carryover_fn, rules, min_long_part):
    """Свод + 12 листов-календарей + список всех заявок. Возвращает bytes .xlsx."""
    wb = Workbook()
    names = [s["name"] for s in staff]

    # --- Свод ---
    ws = wb.active
    ws.title = "Свод"
    ws.append(["ФИО", "Остаток за прошлый год", "Положено всего", "Запланировано", "Одобрено", "На рассмотрении",
               "Осталось", "Есть часть от %s" % n_days(min_long_part), "Отпуска"])
    for c in ws[1]:
        c.font = BOLD
    for name in names:
        s = person_summary(name, records, entitlement_fn(name), min_long_part)
        periods = "; ".join(
            "%s (%s, %s)" % (fmt_period(r["start"], r["end"]), n_days(r["days"]), STATUS_LABELS[r["status"]].lower())
            for r in sorted(records, key=lambda r: r["start"]) if r["name"] == name
        )
        ws.append([name, carryover_fn(name), s["entitlement"], s["planned"], s["approved"], s["pending"],
                   s["entitlement"] - s["planned"], "да" if s["has_long"] else "нет", periods])
        row = ws.max_row
        if s["planned"] < s["entitlement"]:
            ws.cell(row, 7).fill = FILL_TODO
        if not s["has_long"]:
            ws.cell(row, 8).fill = FILL_BAD
    ws.column_dimensions["A"].width = 36
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 14
    ws.column_dimensions["I"].width = 90
    ws.append([])
    ws.append(["Обозначения в календарях:"])
    for label, fill in [("Одобрено", FILL_APPROVED), ("На рассмотрении", FILL_PENDING),
                        ("Праздник", FILL_HOLIDAY), ("Выходной", FILL_WEEKEND),
                        ("Нарушение среди одобренных: пересечение основных дольше %s "
                         "или больше %s в отпуске" % (n_days(rules.max_overlap), n_doctors(rules.max_concurrent)),
                         FILL_BAD)]:
        ws.append([label])
        ws.cell(ws.max_row, 1).fill = fill

    # --- Календари по месяцам ---
    by_day = {}
    for r in active(records):
        for d in daterange(r["start"], r["end"]):
            by_day[(r["name"], d)] = r["status"]
    team = team_calendar(records, rules.main_names)
    bad = bad_days(records, rules)
    for m in range(1, 13):
        ws = wb.create_sheet(MONTH_NAMES[m - 1])
        n = calmod.monthrange(year, m)[1]
        dates = [datetime.date(year, m, d) for d in range(1, n + 1)]
        ws.append(["ФИО"] + [d.day for d in dates])
        ws.append([""] + [DOW_RU[d.weekday()] for d in dates])
        for col, d in enumerate(dates, start=2):
            for row in (1, 2):
                c = ws.cell(row, col)
                c.font = BOLD
                c.alignment = CENTER
                if is_holiday(d):
                    c.fill = FILL_HOLIDAY
                elif d.weekday() >= 5:
                    c.fill = FILL_WEEKEND
        for name in names:
            ws.append([name] + ["От" if (name, d) in by_day else "" for d in dates])
            row = ws.max_row
            for col, d in enumerate(dates, start=2):
                st = by_day.get((name, d))
                c = ws.cell(row, col)
                c.alignment = CENTER
                if st == APPROVED:
                    c.fill = FILL_APPROVED
                elif st == PENDING:
                    c.fill = FILL_PENDING
        for label, key in (("Основных: одобрено", "approved"), ("Основных: на рассмотрении", "pending")):
            ws.append([label] + [len(team.get(d, {}).get(key, [])) or "" for d in dates])
            row = ws.max_row
            ws.cell(row, 1).font = BOLD
            for col, d in enumerate(dates, start=2):
                c = ws.cell(row, col)
                c.alignment = CENTER
                c.font = BOLD
                if key == "approved" and d in bad:
                    c.fill = FILL_BAD
        ws.column_dimensions["A"].width = 36
        for col in range(2, n + 2):
            ws.column_dimensions[get_column_letter(col)].width = 4.5
        ws.freeze_panes = "B3"

    # --- Все заявки ---
    ws = wb.create_sheet("Заявки")
    ws.append(["ФИО", "Начало", "Конец", "Дней", "Статус", "Комментарий сотрудника",
               "Комментарий заведующего", "Подано", "Обновлено"])
    for c in ws[1]:
        c.font = BOLD
    for r in sorted(records, key=lambda r: (names.index(r["name"]) if r["name"] in names else 999, r["start"])):
        ws.append([r["name"], r["start"], r["end"], r["days"], STATUS_LABELS[r["status"]],
                   r["comment"], r["manager_comment"], r["created"], r["updated"]])
        ws.cell(ws.max_row, 2).number_format = "DD.MM.YYYY"
        ws.cell(ws.max_row, 3).number_format = "DD.MM.YYYY"
    for col, w in zip("ABCDEFGHI", [36, 12, 12, 7, 17, 40, 40, 20, 20]):
        ws.column_dimensions[col].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
